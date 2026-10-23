"""Báo cáo chi tiết từ source, ảnh giao diện và output kiểm tra thực tế.

Chạy bằng runtime Python có python-docx và Pillow. Không đọc .env vào báo cáo.
"""
from pathlib import Path
import json, datetime, textwrap, subprocess
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'docs';E=ROOT/'minh-chung';D=E/'chi-tiet'
V=json.loads((E/'verification.json').read_text())
R=json.loads((E/'reproducibility.json').read_text())
C=json.loads((D/'command-evidence.json').read_text())
CHECK=datetime.datetime.fromisoformat(V['time']).strftime('%H:%M:%S ngày %d/%m/%Y')
CAPTURE=datetime.datetime.fromisoformat(C['time']).strftime('%H:%M:%S ngày %d/%m/%Y')
REPO='https://github.com/longphan632000/DTC245220017_De29-Recipe'
doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.left_margin=Cm(2.2);sec.right_margin=Cm(2.0);sec.top_margin=Cm(2.0);sec.bottom_margin=Cm(1.8)
sec.header_distance=Cm(.8);sec.footer_distance=Cm(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption','List Number','List Bullet']:
    style=doc.styles[name];style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0)
    rf=style.element.get_or_add_rPr().get_or_add_rFonts()
    for a in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+a),'Times New Roman')
    for a in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        rf.attrib.pop(qn('w:'+a),None)
normal=doc.styles['Normal'];normal.font.size=Pt(12)
normal.paragraph_format.line_spacing=1.15;normal.paragraph_format.space_after=Pt(6)
for name,size in [('Title',20),('Heading 1',16),('Heading 2',12.5),('Heading 3',12)]:
    st=doc.styles[name];st.font.size=Pt(size);st.font.bold=True
    st.paragraph_format.space_before=Pt(5);st.paragraph_format.space_after=Pt(8)
doc.styles['Caption'].font.size=Pt(10);doc.styles['Caption'].font.italic=True
doc.styles['Caption'].paragraph_format.space_after=Pt(7)
header=sec.header.paragraphs[0];header.text='TRIỂN KHAI VÀ QUẢN TRỊ HỆ THỐNG PHẦN MỀM  |  ĐỀ 29 RECIPE'
for r in header.runs:r.font.size=Pt(8)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('Trang ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title='Báo cáo chi tiết triển khai Website Công thức Nấu ăn Recipe'
doc.core_properties.author='DTC245220017'
doc.core_properties.subject='Quy trình triển khai cấu hình kiểm thử và minh chứng thực tế Đề 29'
pages=[];figures=[];chapter_rows=[]

def p(text):
    x=doc.add_paragraph(text);x.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;return x
def h(text):doc.add_heading(text,level=2)
def page(title,chapter=False):
    if pages:doc.add_page_break()
    pages.append(title)
    doc.add_heading(title,level=1)
    if chapter:chapter_rows.append((title,len(pages)))
def steps(items):
    for i,text in enumerate(items,1):
        x=p(str(i)+'. '+text);x.paragraph_format.space_after=Pt(4)
def code(text,size=9):
    for line in text.strip('\n').splitlines():
        # Chỉ ngắt hiển thị; các lệnh dài trong văn bản được xuống dòng rõ ràng.
        for line2 in textwrap.wrap(line,91,break_long_words=True,break_on_hyphens=False,replace_whitespace=False,drop_whitespace=False) or ['']:
            x=doc.add_paragraph();x.paragraph_format.line_spacing=1.0;x.paragraph_format.space_after=Pt(1)
            run=x.add_run(line2);run.font.name='Liberation Mono';run.font.size=Pt(size)
    x=doc.add_paragraph();x.paragraph_format.space_after=Pt(1);x.paragraph_format.line_spacing=Pt(2);x.add_run().font.size=Pt(2)
def table(headers,rows,widths=None,size=10.5):
    if widths is None:widths=[6.5/len(headers)]*len(headers)
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,w in enumerate(widths):t.columns[i].width=Inches(w)
    for i,x in enumerate(headers):t.rows[0].cells[i].text=str(x)
    for values in rows:
        cells=t.add_row().cells
        for i,x in enumerate(values):cells[i].text=str(x)
    for ri,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if not ri:pr.append(OxmlElement('w:tblHeader'))
        for ci,cell in enumerate(row.cells):
            cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props=cell._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                q=OxmlElement('w:'+edge);q.set(qn('w:val'),'single');q.set(qn('w:sz'),'4');q.set(qn('w:color'),'D9D9D9');b.append(q)
            props.append(b);m=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                q=OxmlElement('w:'+edge);q.set(qn('w:w'),'75');q.set(qn('w:type'),'dxa');m.append(q)
            props.append(m);shade=OxmlElement('w:shd')
            shade.set(qn('w:fill'),'DCE6F1' if ri==0 else ('F5F7F8' if ri%2==0 else 'FFFFFF'));props.append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.line_spacing=1.05;para.paragraph_format.space_after=Pt(2)
                for r in para.runs:r.font.size=Pt(size);r.bold=ri==0
    x=doc.add_paragraph();x.paragraph_format.space_after=Pt(2);x.paragraph_format.line_spacing=Pt(2);x.add_run().font.size=Pt(2)
    return t
def fig(name,caption,maxheight=3.55):
    path=E/name
    if not path.exists():raise FileNotFoundError(path)
    with Image.open(path) as im:ratio=im.width/im.height
    width=min(6.4,maxheight*ratio)
    x=doc.add_paragraph();x.alignment=WD_ALIGN_PARAGRAPH.CENTER
    x.paragraph_format.space_after=Pt(2);x.paragraph_format.keep_with_next=True
    shape=x.add_run().add_picture(str(path),width=Inches(width))
    shape._inline.docPr.set('descr',caption)
    number=len(figures)+1
    c=doc.add_paragraph('Hình '+str(number)+' '+caption,style='Caption');c.alignment=WD_ALIGN_PARAGRAPH.CENTER
    figures.append({'number':number,'caption':caption,'file':str(path.relative_to(ROOT)),'planned_page':len(pages)})
def ui(n,tail,caption,maxheight=3.45):fig('chi-tiet/UI_'+str(n).zfill(2)+'_'+tail+'.png',caption,maxheight)
def cli(key,caption,maxheight=4.7):fig('chi-tiet/'+key+'.png',caption,maxheight)

def flow_diagram():
    im=Image.new('RGB',(1600,880),'white');d=ImageDraw.Draw(im)
    f=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',27)
    b=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf',29)
    def box(x,y,w,hh,title,lines):
        d.rounded_rectangle((x,y,x+w,y+hh),14,fill='#f1f4f2',outline='#789186',width=2)
        d.text((x+18,y+15),title,font=b,fill='black')
        for i,line in enumerate(lines):d.text((x+18,y+62+i*34),line,font=f,fill='black')
    def arrow(x1,y1,x2,y2,label):
        import math
        d.line((x1,y1,x2,y2),fill='#355a49',width=4);a=math.atan2(y2-y1,x2-x1)
        d.polygon([(x2,y2),(x2-18*math.cos(a-.45),y2-18*math.sin(a-.45)),(x2-18*math.cos(a+.45),y2-18*math.sin(a+.45))],fill='#355a49')
        if label:d.text(((x1+x2)/2-70,(y1+y2)/2-35),label,font=f,fill='#355a49')
    d.text((45,25),'Hai luồng dữ liệu quan sát trong hệ thống Recipe',font=b,fill='black')
    box(40,120,420,190,'Nguồn metrics',['cAdvisor CPU RAM mạng','Nginx Exporter','MySQL Exporter'])
    box(650,120,370,190,'Prometheus',['Chủ động scrape mỗi 15s','Lưu mẫu theo thời gian'])
    box(1200,310,350,190,'Grafana',['Datasource Prometheus','Datasource Loki','Dashboard và Explore'])
    box(40,530,420,170,'Nguồn log',['File Nginx PHP MySQL','Named volume log'])
    box(650,530,370,170,'Promtail và Loki',['Promtail đọc và push log','Loki lưu và xử lý LogQL'])
    arrow(460,210,650,210,'Metrics');arrow(1020,220,1200,350,'PromQL')
    arrow(460,615,650,615,'Dòng log');arrow(1020,615,1200,460,'LogQL')
    d.text((45,765),'Mũi tên mô tả chiều dữ liệu trả về cho bên sử dụng.',font=f,fill='black')
    d.text((45,810),'Prometheus pull metrics; Promtail push log; Grafana gửi truy vấn tới hai kho.',font=f,fill='black')
    im.save(D/'So_do_metrics_logs.png')
flow_diagram()

pages.append('Nội dung báo cáo')
doc.add_paragraph('Báo cáo chi tiết triển khai Website Công thức Nấu ăn Recipe',style='Title')
p('Báo cáo mô tả quá trình xây dựng, triển khai và kiểm tra Đề 29 trên Ubuntu bằng Docker Compose. Nội dung đi từ yêu cầu và cấu hình nguồn đến thao tác sử dụng, kết quả kiểm thử và cách vận hành. Phần bìa được bổ sung riêng khi hoàn thiện hồ sơ nộp.')
h('Nội dung báo cáo')
toc=table(['Phần nội dung','Trang'],[],[5.9,.6],10.5)
p('Hình giao diện được chụp từ các dịch vụ đang chạy. Hình kết quả lệnh là bản định dạng lại output thực tế, có file TXT và JSON nguồn đi kèm để đối chiếu; đây không phải ảnh chụp cửa sổ Terminal. Các sơ đồ được xây dựng từ cấu hình triển khai.')
p('Các thao tác tạo và sửa dữ liệu minh họa dùng bản ghi tạm riêng, sau đó được dọn đúng ID. Sáu công thức mẫu và các dữ liệu có sẵn khác được giữ nguyên.')

page('1 Mục tiêu và tiêu chí thực hiện',True)
p('Đề 29 yêu cầu triển khai website chia sẻ công thức nấu ăn, có món ăn, nguyên liệu và danh mục; sử dụng MySQL và phpMyAdmin. Các thành phần vận hành gồm GitHub, Nginx reverse proxy, Prometheus, Grafana, Loki, Promtail và các biện pháp hardening. Báo cáo cung cấp đường dẫn source, thao tác tái lập và minh chứng cho từng nhóm yêu cầu.')
table(['Tiêu chí','Sản phẩm đã triển khai','Minh chứng chính'],[
('GitHub 1,5 điểm','Source, cấu hình, README và commit có ý nghĩa','Repository và lịch sử commit'),
('Ứng dụng và DB 1,5 điểm','PHP Recipe, MySQL, phpMyAdmin','Giao diện, SQL, kiểm thử CRUD'),
('Reverse proxy 1,5 điểm','Nginx chuyển tiếp website và DB tool','Cấu hình, nginx -t, header HTTP'),
('Giám sát 1,5 điểm','Metrics container, web và database','Targets, PromQL, dashboard'),
('Log tập trung 1,5 điểm','Ba nguồn log, Loki và Promtail','Bốn truy vấn LogQL có dữ liệu'),
('Hardening 1,5 điểm','Non-root, mạng nội bộ, giảm quyền','Docker inspect, DB grants, HTTP'),
('Tổng thể 1 điểm','Docker Compose, báo cáo và kịch bản demo','Kiểm thử, ảnh và hướng dẫn')],[1.5,2.5,2.5],10)
p('Phương án chọn ứng dụng PHP tự xây theo cách tổ chức thư mục của repository mẫu Blog CMS. Mã Recipe và dữ liệu công thức được xây dựng cho đề tài này. WordPress đã cài thử trước đó không còn là ứng dụng phục vụ chính; các bảng và volume cũ được giữ để bảo toàn dữ liệu.')
p('Kết quả tích hợp mới nhất ghi nhận '+str(V['passed'])+'/'+str(V['total'])+' kiểm tra đạt lúc '+CHECK+'. Kiểm tra cài mới độc lập từ GitHub đã đạt '+str(R['passed'])+'/'+str(R['total'])+' ở commit '+R['commit'][:7]+'. Hai mốc kiểm tra được phân biệt trong phần kiểm thử.')

page('2 Kiến trúc và luồng truy cập',True)
fig('So_do_kien_truc.png','Kiến trúc ứng dụng và các thành phần quan sát',3.65)
p('Nginx là điểm nhận HTTP của website tại 127.0.0.1:8088. Request được chuyển tới Apache PHP tại app:8080. PHP dùng PDO để đọc hoặc cập nhật MySQL qua hostname db. Cổng 8081 của Nginx chuyển tới phpMyAdmin; database không mở cổng trực tiếp ra host.')
table(['Lớp','Thành phần','Vai trò'],[('Giao diện','PHP và CSS SVG','Danh sách, tìm kiếm, chi tiết và quản trị'),('Dữ liệu','MySQL và phpMyAdmin','Lưu công thức, kiểm tra và quản lý dữ liệu'),('Giám sát','Prometheus Grafana exporters','Theo dõi trạng thái và tài nguyên'),('Nhật ký','Promtail Loki Grafana','Gom log và tìm sự kiện')],[1.1,2.2,3.2])
p('Dữ liệu nghiệp vụ nằm trong MySQL. Metrics nằm trong Prometheus; log tập trung nằm trong Loki. Grafana truy vấn hai kho quan sát này để hiển thị. Việc tách vai trò giúp tìm được cả thời điểm bất thường từ biểu đồ và các sự kiện liên quan từ log.')

page('3 Chuẩn bị môi trường Ubuntu',True)
steps(['Mở Terminal tại thư mục dự án trên Ubuntu. Kiểm tra Docker Engine, Compose, Git, GitHub CLI và Python 3.','Xác nhận Docker daemon hoạt động và tài khoản hiện tại có thể chạy docker compose.','Kiểm tra các cổng dự kiến sử dụng để tránh trùng với ứng dụng khác trên máy.'])
code('cd -- "/home/longphan/DTC245220017-De 29"\ndocker --version\ndocker compose version\ngit --version\ngh --version\npython3 --version')
cli('CLI_01_Moi_truong','Phiên bản công cụ được kiểm tra trực tiếp trên Ubuntu',2.8)
table(['Cổng trên localhost','Chức năng'],[('8088','Website và trang quản trị Recipe'),('8081','phpMyAdmin qua Nginx'),('3000','Grafana'),('9090','Prometheus'),('3100','Loki API')],[1.8,4.7])
p('Các URL localhost phải được mở trong Ubuntu đang chạy Docker. Khi truy cập từ máy khác, dùng SSH tunnel tới Ubuntu. Giữ các UI quản trị trên loopback giúp giảm phạm vi truy cập trong môi trường lab.')

page('4 Quản lý mã nguồn trên GitHub',True)
p('Repository của sinh viên là '+REPO+'. Tên repo chứa mã DTC245220017 và chủ đề Recipe. Repository công khai chứa mã nguồn, cấu hình và tài liệu hướng dẫn; các bí mật chạy thật được loại khỏi Git.')
ui(20,'GitHub_repository','Repository GitHub có source và các thư mục cấu hình',3.25)
ui(21,'GitHub_lich_su_commit','Lịch sử bốn mốc commit tại thời điểm chụp minh chứng',3.25)
p('Ba commit kỹ thuật đầu tương ứng triển khai ứng dụng và Nginx, giám sát, log tập trung và hardening. Các commit tài liệu sau bổ sung báo cáo và ảnh minh chứng.')

page('Cấu trúc source và ý nghĩa các commit')
table(['Đường dẫn','Trách nhiệm'],[
('app/src','PHP, template HTML, CSS và sáu hình SVG'),
('app/Dockerfile','Build PHP Apache chạy UID 33 và cổng 8080'),
('mysql/init.sql','Tạo schema và dữ liệu công thức ban đầu'),
('docker/','Dockerfile và cấu hình hỗ trợ MySQL phpMyAdmin Nginx Loki Promtail'),
('nginx/nginx.conf','Reverse proxy, headers, log JSON và giới hạn request'),
('prometheus/','Danh sách target và luật cảnh báo'),
('grafana/','Provisioning datasource và dashboard JSON'),
('loki/ và promtail/','Lưu log tập trung và cấu hình nguồn file'),
('scripts/','Khởi tạo, kiểm thử, sao lưu và thu minh chứng'),
('docs/ và minh-chung/','Báo cáo, hướng dẫn, ảnh, kết quả API và TXT')],[1.7,4.8],10.3)
cli('CLI_02_Git','Các commit kỹ thuật và thời điểm từ git log',3.15)
p('Cách chia thư mục tham khảo repository Khanhslee/DTC245220024_De7-Blog-CMS. Việc tham khảo giới hạn ở khung tổ chức hệ thống; nội dung Recipe, provisioning và bộ kiểm thử được xây cho Đề 29.')
h('Các lệnh quản lý repository')
p('Lần đầu tạo repo dùng git init, commit rồi GitHub CLI. Những lần cập nhật sau chỉ stage đúng thay đổi, commit và push; không tạo lại remote đã tồn tại.')
code('git init -b main\ngit add .\ngit commit -m "feat: deploy Recipe PHP app with MySQL and Nginx"\ngh repo create longphan632000/DTC245220017_De29-Recipe \\\n  --public --source=. --remote=origin --push\ngit push origin main')

page('5 Quy trình khởi tạo bằng Docker Compose',True)
steps(['Clone repository và chuyển vào thư mục source.','Chạy setup.py để tạo bí mật cục bộ, chuẩn bị secret phpMyAdmin và kiểm tra Compose.','Script build image, chạy MySQL đến khi healthy rồi nạp schema và dữ liệu mẫu.','Script giới hạn quyền DB, tạo tài khoản exporter, chạy các dịch vụ còn lại và khởi tạo quản trị viên bằng password hash.','Kiểm tra trạng thái container, HTTP và chức năng bằng verify.py.'])
code('git clone '+REPO+'.git\ncd DTC245220017_De29-Recipe\npython3 scripts/setup.py\ndocker compose ps -a\npython3 scripts/verify.py')
h('Vì sao cần script setup')
p('Ứng dụng chạy bằng tài khoản DB chỉ có quyền DML nên không tự tạo bảng khi khởi động. Setup thực hiện thao tác quản trị cần thiết trong giai đoạn khởi tạo rồi cấp SELECT, INSERT, UPDATE, DELETE cho tài khoản ứng dụng. Trình tự này tách quyền triển khai schema khỏi quyền chạy website.')
p('Nếu .env chưa tồn tại, script tạo mật khẩu ngẫu nhiên riêng cho database, ứng dụng và Grafana. Nếu đã có .env, script giữ giá trị hiện tại. .env đặt quyền 0600, thư mục secrets đặt 0700; các file này không được đưa lên GitHub.')
table(['Bước trong script','Điều kiện hoặc kết quả'],[
('docker compose config --quiet','YAML và phép nội suy biến hợp lệ'),
('docker compose build','Tạo các image tùy chỉnh theo Dockerfile'),
('up --wait db','Database vượt healthcheck trước khi nạp dữ liệu'),
('Nạp mysql/init.sql','Schema idempotent và dữ liệu tiếng Việt utf8mb4'),
('Cấp quyền và khởi tạo admin','Mật khẩu băm, quyền DB theo vai trò'),
('up --wait toàn bộ','Khởi chạy hệ thống bằng một project Compose')],[2.0,4.5])
h('Các thuộc tính Compose của ứng dụng')
code('app:\n  user: "33:33"\n  read_only: true\n  networks: [web_net, db_net]\n  depends_on:\n    db:\n      condition: service_healthy\n  mem_limit: 384m')
p('Đây là phần trích của service app. Dockerfile cài pdo_mysql, chuyển Apache sang cổng 8080, COPY source với chủ sở hữu UID 33 và khai báo USER 33:33. tmpfs cung cấp nơi ghi file tạm cho Apache khi root filesystem chỉ đọc.')

page('Container và volume sau khi triển khai')
cli('CLI_03_Containers','Mười một dịch vụ của project Recipe đang chạy',3.25)
cli('CLI_12_Volumes','Các named volume được Docker quản lý',2.5)
p('App, MySQL và Nginx có healthcheck kiểm tra điều kiện phục vụ riêng. Một số dịch vụ hiển thị running mà không có Docker healthcheck; mức hoạt động của chúng được kiểm chứng thêm bằng HTTP, datasource health hoặc metrics target trong bộ kiểm thử.')
p('Named volume giữ database, dữ liệu Grafana, metrics, log và vị trí đọc Promtail qua các lần tạo lại container. Dừng bằng docker compose stop không xóa dữ liệu. Lệnh down -v sẽ xóa volume và không dùng khi cần giữ kết quả.')

page('Phân tách mạng và cổng truy cập')
cli('CLI_05_Networks','Các mạng Docker và thành viên tại thời điểm kiểm tra',3.1)
cli('CLI_06_Ports_secrets','Cổng chỉ bind localhost và quyền file bí mật',2.8)
p('db_net và monitoring_net là mạng internal. web_net nối Nginx với app và phpMyAdmin. management_net nối các UI giám sát với cơ chế publish cổng trên host. Thành viên của mỗi mạng được giới hạn theo nhu cầu kết nối; app không được nối trực tiếp với monitoring_net.')
p('Network internal giúp giới hạn đường đi mạng nhưng không thay thế xác thực dịch vụ. Các cổng publish đều dùng HostIp 127.0.0.1; MySQL và các exporter không có cổng host.')

page('6 Thiết kế cơ sở dữ liệu Recipe',True)
fig('So_do_du_lieu.png','Quan hệ giữa năm bảng dữ liệu của ứng dụng Recipe',3.4)
table(['Bảng','Dữ liệu chính','Quan hệ'],[
('categories','Tên danh mục','Một danh mục có nhiều món'),
('recipes','Tên, mô tả, thời gian, khẩu phần, độ khó','category_id tham chiếu categories'),
('ingredients','Tên nguyên liệu và thứ tự','recipe_id tham chiếu recipes'),
('steps','Hướng dẫn và thứ tự bước','recipe_id tham chiếu recipes'),
('users','Tên đăng nhập, hash mật khẩu, vai trò','Tài khoản quản trị nội dung')],[1.1,2.6,2.8],10.2)
p('Schema dùng InnoDB và utf8mb4. Khóa ngoại giữ liên kết giữa món ăn và các dòng nguyên liệu, cách làm. Ràng buộc thứ tự theo từng công thức giúp kết quả được trình bày ổn định. is_deleted phục vụ lưu trữ và khôi phục thay vì xóa ngay công thức.')
h('Trích khai báo bảng nguyên liệu')
code('CREATE TABLE IF NOT EXISTS ingredients (\n id INT AUTO_INCREMENT PRIMARY KEY, recipe_id INT NOT NULL,\n name VARCHAR(500) NOT NULL, position SMALLINT UNSIGNED NOT NULL,\n UNIQUE KEY uq_ingredient(recipe_id,position),\n FOREIGN KEY(recipe_id) REFERENCES recipes(id) ON DELETE CASCADE\n) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;')

page('Khởi tạo dữ liệu và kiểm tra phpMyAdmin')
steps(['Trong setup, nạp mysql/init.sql sau khi MySQL healthy.','Mở http://127.0.0.1:8081, đăng nhập bằng MYSQL_USER và MYSQL_PASSWORD trong .env.','Chọn recipe_db, đối chiếu số dòng và engine của các bảng ứng dụng.'])
cli('CLI_07_Database','Số dòng dữ liệu và quyền DML thực tế của recipe_user',2.55)
ui(17,'phpMyAdmin_cau_truc','Năm bảng Recipe và các bảng WordPress cũ được giữ trên máy',3.25)
p('Dữ liệu mẫu có 4 danh mục, 6 món, 36 nguyên liệu, 27 bước và 1 quản trị viên. Các bảng wp_ là dữ liệu từ lần cài thử trước, không được ứng dụng PHP sử dụng. Cài mới từ GitHub chỉ tạo đúng năm bảng ứng dụng.')

page('Đối chiếu dữ liệu bằng truy vấn SQL')
p('Tại phpMyAdmin, chọn recipe_db rồi mở tab SQL. Câu truy vấn sau nối món ăn với danh mục, tính tổng thời gian và chỉ lấy công thức công khai.')
code('SELECT r.id, r.title, c.name AS danh_muc,\n       r.prep_minutes + r.cook_minutes AS tong_phut,\n       r.servings\nFROM recipes r\nJOIN categories c ON c.id = r.category_id\nWHERE r.is_deleted = 0\nORDER BY r.id;')
ui(19,'SQL_doi_chieu_cong_thuc','Truy vấn JOIN trả sáu công thức cùng danh mục thời gian khẩu phần',3.65)
p('Kết quả phở gà 70 phút và 4 người được đối chiếu với trang chi tiết của website. Cách kiểm tra này xác nhận giao diện lấy dữ liệu từ database, đồng thời minh họa quan hệ recipes và categories.')
p('Thông báo phpMyAdmin về việc không có cột duy nhất để sửa trực tiếp lưới kết quả JOIN không phải lỗi thực thi SELECT. Truy vấn vẫn trả dữ liệu; chỉnh sửa nghiệp vụ thực hiện qua trang quản trị ứng dụng.')

page('7 Giao diện website và danh sách công thức',True)
p('Người dùng mở http://127.0.0.1:8088 qua Nginx. Trang chủ giới thiệu website; phần danh sách hiển thị thẻ món ăn, danh mục, thời gian, khẩu phần và độ khó. Các hình SVG nằm trong source nên giao diện không phụ thuộc dịch vụ ảnh bên ngoài.')
fig('Hinh_01_Trang_chu_Recipe.png','Trang chủ Recipe được truy cập qua cổng Nginx',3.2)
fig('Hinh_02_Danh_sach_cong_thuc.png','Các thẻ công thức và bộ lọc danh mục',3.2)
p('Số công thức lấy bằng COUNT trong MySQL. Nội dung danh sách được truy vấn theo điều kiện is_deleted=0 để loại món đã lưu trữ.')

page('Thao tác tìm nguyên liệu và lọc danh mục')
p('Nhập me chín vào ô tìm kiếm rồi chạy tìm. Ứng dụng tìm đồng thời trong tên món, mô tả và bảng nguyên liệu; kết quả là một công thức canh chua. Chọn Rau và salad để kiểm tra bộ lọc danh mục.')
ui(11,'Tim_kiem_nguyen_lieu','Từ khóa me chín trả đúng món canh chua',3.35)
ui(12,'Loc_danh_muc','Danh mục Rau và salad trả một công thức salad',3.35)
p('Câu SQL dùng tham số PDO và điều kiện EXISTS trên ingredients. Bộ lọc danh mục thêm điều kiện category_id; từ khóa được escape khi đưa lại lên HTML.')

page('Thao tác xem chi tiết và chuẩn bị nguyên liệu')
p('Mở Phở gà thơm vị nhà từ danh sách. Kiểm tra tổng thời gian 70 phút, khẩu phần 4 người, sáu nguyên liệu và năm bước. Các dòng nguyên liệu có checkbox để đánh dấu trong lúc chuẩn bị.')
ui(13,'Chi_tiet_mon_an','Thông tin món ăn thời gian và khẩu phần từ database',3.35)
ui(14,'Nguyen_lieu_cac_buoc','Nguyên liệu và hướng dẫn được sắp theo thứ tự',3.35)
p('Nút In công thức dùng CSS dành cho bản in. Checkbox chỉ hỗ trợ thao tác trên giao diện, không cập nhật database.')

page('Đăng nhập và mở trang quản trị')
steps(['Mở đường dẫn ?page=login. Nhập APP_ADMIN_USER và APP_ADMIN_PASSWORD từ .env.','Khi đăng nhập thành công, hệ thống đổi session ID và chuyển đến trang quản trị.','Trang quản trị cung cấp thêm món, sửa món, danh sách đang hiển thị và danh sách lưu trữ.'])
ui(1,'Dang_nhap','Biểu mẫu đăng nhập quản trị trước khi nhập mật khẩu',3.15)
ui(2,'Quan_tri','Trang quản trị hiển thị danh sách công thức hiện có',3.15)
p('Người chưa có session hợp lệ được chuyển về đăng nhập khi mở ?page=admin hoặc ?page=edit. Tài khoản quản trị được tạo bằng hash mật khẩu trong bảng users.')

page('Thao tác thêm công thức mới')
p('Từ trang quản trị chọn Thêm công thức. Bản ghi minh họa dùng tên Cơm gà minh chứng báo cáo, sơ chế 10 phút, nấu 25 phút, 2 người. Nguyên liệu và cách làm được nhập mỗi dòng một mục.')
ui(3,'Them_cong_thuc_thong_tin','Nhập tên mô tả danh mục thời gian và khẩu phần',3.35)
ui(4,'Them_nguyen_lieu_cach_nau','Nhập nguyên liệu cách làm và nút lưu công thức',3.35)
p('Form kiểm tra trường bắt buộc, danh mục hợp lệ, số người từ 1 đến 50 và mỗi thời gian không quá 600 phút. Lưu dữ liệu món và các dòng con trong cùng transaction.')

page('Xác nhận tạo mới và mở biểu mẫu chỉnh sửa')
ui(5,'Them_thanh_cong','Thông báo lưu thành công và món mới có tổng thời gian 35 phút',3.4)
ui(6,'Chinh_sua_cong_thuc','Biểu mẫu chỉnh sửa cùng bản ghi minh họa',3.4)
p('Bản ghi minh họa có ID 14. Sau khi tạo, chọn Sửa ở đúng dòng này. Đổi tên thành Cơm gà minh chứng đã chỉnh sửa và thời gian nấu thành 30 phút. Khẩu phần, nguyên liệu và các bước được giữ để đối chiếu.')
p('Trong mã PHP, thao tác sửa UPDATE bảng recipes và thay các dòng nguyên liệu, cách làm trong một transaction. Nếu có lỗi database, dữ liệu không được cập nhật dở dang.')

page('Kiểm tra cập nhật và lưu trữ công thức')
ui(7,'Sua_thanh_cong','Tên món được cập nhật và tổng thời gian đổi thành 40 phút',3.4)
ui(8,'Luu_tru_cong_thuc','Bản ghi chuyển sang danh sách đã lưu trữ',3.4)
p('Sau khi lưu, kiểm tra tên mới và tổng thời gian 10 + 30 = 40 phút. Chọn Lưu trữ rồi mở tab Đã lưu trữ. Thao tác này đặt is_deleted=1 và giữ dữ liệu trong MySQL để có thể khôi phục.')

page('Đối chiếu trạng thái ẩn và khôi phục')
ui(9,'Mon_luu_tru_404','Công thức đã lưu trữ không còn được phục vụ ở trang chi tiết',3.25)
ui(10,'Khoi_phuc_cong_thuc','Cùng công thức hiển thị lại sau khi khôi phục',3.25)
p('Khi món đã lưu trữ, URL ?recipe=14 trả trang không tìm thấy và HTTP 404. Trong danh sách lưu trữ, chọn Khôi phục để đặt is_deleted=0; tải lại URL và kiểm tra món trở lại với tên mới, 40 phút, 2 người.')
p('Ảnh ghi nhận trạng thái khi thao tác. Sau chu trình minh chứng, bản ghi tạm ID 14 được dọn theo đúng ID và tên; database trở về sáu công thức mẫu. Các món gốc không bị dùng để thử xóa hoặc chỉnh sửa.')

page('Cơ chế bảo vệ trong mã ứng dụng')
h('Truy vấn và hiển thị dữ liệu')
code("$s = db()->prepare('SELECT id,username,password_hash FROM users WHERE username=? AND role=?');\n$s->execute([$name, 'admin']);\nhtmlspecialchars($value, ENT_QUOTES | ENT_SUBSTITUTE, 'UTF-8');")
p('Prepared statement tách câu SQL khỏi giá trị đầu vào. Chuỗi có dấu nháy hoặc nội dung giống SQL vẫn được xử lý như dữ liệu. Escape HTML ngăn nội dung tên món và mô tả trở thành thẻ hoặc script khi xuất ra giao diện.')
h('Session mật khẩu và CSRF')
code("password_verify($password, $user['password_hash']);\nsession_regenerate_id(true);\nhash_equals(csrf(), $_POST['csrf']);")
p('Mật khẩu lưu bằng password_hash, không lưu dạng rõ trong bảng users. Đăng nhập thành công đổi session ID. Các POST như đăng nhập, lưu, lưu trữ và khôi phục đều qua kiểm tra CSRF token; token thiếu hoặc sai nhận HTTP 403. Cookie session đặt HttpOnly và SameSite=Strict.')
table(['Tình huống kiểm tra','Kết quả thực tế trong verify.py'],[
('Chưa đăng nhập mở quản trị','Chuyển về trang đăng nhập'),
('POST thiếu CSRF','HTTP 403'),
('Sai mật khẩu','HTTP 401 và sự kiện login_failed'),
('Tìm chuỗi SQL injection','Được xử lý như từ khóa'),
('Nhập nội dung chứa HTML','Nội dung được HTML escape'),
('Đăng nhập thành công','Session ID thay đổi')],[2.5,4.0])
p('Audit log ghi tên sự kiện và ID bản ghi, không ghi mật khẩu, cookie hay CSRF token. Bộ kiểm thử tự tạo dữ liệu riêng và dọn bản ghi trong khối finally để giảm ảnh hưởng tới dữ liệu mẫu.')

page('8 Cấu hình Nginx reverse proxy',True)
p('Nginx dùng hai server block: cổng nội bộ 8080 phục vụ website và 8081 phục vụ phpMyAdmin. Host publish lần lượt 127.0.0.1:8088 và 127.0.0.1:8081. Apache app và phpMyAdmin chỉ có cổng nội bộ Docker.')
code('resolver 127.0.0.11 valid=10s ipv6=off;\nupstream recipe_backend {\n    zone recipe_backend 64k;\n    server app:8080 resolve;\n}\nlocation / {\n    proxy_set_header Host $http_host;\n    proxy_set_header X-Forwarded-For $proxy_add_x_forwarded_for;\n    proxy_pass http://recipe_backend;\n}')
p('Resolver Docker và tùy chọn resolve giúp Nginx cập nhật địa chỉ app khi container được tạo lại. Các header chuyển tiếp giữ host và thông tin proxy cần thiết cho ứng dụng.')
steps(['Sửa nginx/nginx.conf trong source.','Build lại image Nginx khi thay đổi cấu hình COPY vào image.','Chạy nginx -t trong container để xác nhận cú pháp và kiểm tra HTTP qua cổng proxy.'])
code('docker compose build nginx\ndocker compose up -d --no-deps nginx\ndocker compose exec -T nginx nginx -t')
cli('CLI_10_Nginx_config_test','Kết quả kiểm tra cú pháp cấu hình Nginx',1.55)

page('Security headers và chặn đường dẫn nhạy cảm')
cli('CLI_09_HTTP_headers','HTTP 200 403 404 và security headers từ Nginx',3.55)
table(['Header','Tác dụng'],[
('X-Content-Type-Options','nosniff hạn chế đoán sai kiểu nội dung'),
('X-Frame-Options','SAMEORIGIN hạn chế nhúng trang từ nguồn khác'),
('Referrer-Policy','Giới hạn thông tin referrer gửi tới nguồn khác'),
('Permissions-Policy','Tắt camera microphone và geolocation không cần dùng'),
('Content-Security-Policy','Giới hạn các nguồn script style ảnh kết nối form')],[2.1,4.4],10.2)
p('Nginx ẩn số phiên bản, chặn file cấu hình như /.env và giới hạn request đăng nhập theo địa chỉ client. Bản lab chọn security headers theo phương án đề bài cho phép; không mô tả website localhost HTTP như một triển khai HTTPS công khai.')
p('Access log dùng JSON gồm time, method, uri, status, bytes, request_time và upstream. Trường uri không kèm query string để giảm việc lưu dữ liệu người dùng vào log.')

page('9 Luồng giám sát bằng Prometheus và Grafana',True)
fig('chi-tiet/So_do_metrics_logs.png','Luồng metrics và log cùng được truy vấn qua Grafana',3.65)
p('Prometheus chủ động gọi endpoint metrics theo scrape_interval 15 giây và lưu mẫu kèm timestamp. cAdvisor cung cấp CPU, RAM và mạng container; Nginx Exporter đọc stub_status; MySQL Exporter dùng tài khoản recipe_exporter để đọc số liệu database.')
p('Grafana không tự thu CPU từ Docker. Mỗi panel chỉ định datasource và biểu thức PromQL; Grafana gửi yêu cầu tới Prometheus cho khoảng thời gian được chọn rồi dùng kết quả để vẽ biểu đồ. Refresh 10 giây của dashboard không làm chu kỳ scrape 15 giây thay đổi.')
code('global:\n  scrape_interval: 15s\n  evaluation_interval: 15s\nscrape_configs:\n  - job_name: nginx\n    static_configs: [{targets: ["nginx-exporter:9113"]}]\n  - job_name: mysql\n    static_configs: [{targets: ["mysql-exporter:9104"]}]')

page('Kiểm tra target và endpoint metrics')
ui(15,'Prometheus_targets','Sáu nhóm target Prometheus đều có trạng thái UP',3.25)
ui(16,'Target_MySQL_chi_tiet','Endpoint MySQL Exporter và nhãn job instance',3.25)
p('Mở Prometheus tại :9090/targets. Sáu target là prometheus, cadvisor, nginx, mysql, loki và promtail. Mở nhóm mysql để xem endpoint mysql-exporter:9104/metrics, lần scrape và trạng thái.')
p('Sáu target không đồng nghĩa sáu container: một endpoint cAdvisor cung cấp metrics của mười một container. up=1 chứng minh scrape endpoint thành công; kiểm tra mysql_up và nginx_up mới xác minh exporter đọc được dịch vụ phía sau.')

page('Truy vấn PromQL và đối chiếu API')
ui(24,'PromQL_up','Câu up trả sáu chuỗi dữ liệu có giá trị một',3.25)
cli('CLI_13_Targets_API','API Prometheus trả endpoint và health của sáu target',2.6)
code('up\nsum(up)\nup{job="mysql"}\nmysql_up\nnginx_up')
p('up trả một chuỗi cho mỗi target với nhãn job và instance. sum(up) cộng các target scrape thành công. Nếu up của mysql bằng 1 nhưng mysql_up bằng 0, exporter vẫn trả metrics nhưng không kết nối được MySQL. Tình huống này cần kiểm tra database, tài khoản exporter và đường mạng db_net.')

page('Provisioning nguồn dữ liệu Grafana')
ui(25,'Grafana_datasources','Hai datasource Loki và Prometheus được khai báo sẵn',3.3)
code('datasources:\n  - name: Prometheus\n    uid: prometheus\n    type: prometheus\n    access: proxy\n    url: http://prometheus:9090\n    isDefault: true\n  - name: Loki\n    uid: loki\n    type: loki\n    access: proxy\n    url: http://loki:3100')
p('Datasource nằm trong grafana/provisioning/datasources. Các URL trên dùng tên dịch vụ Docker, khác với URL localhost người dùng mở trên Ubuntu. Dashboard JSON và cấu hình provider được mount chỉ đọc để nạp tự động khi Grafana khởi động.')
p('Đăng nhập Grafana bằng GRAFANA_ADMIN_USER và GRAFANA_ADMIN_PASSWORD trong .env. Bộ kiểm thử gọi health API của cả hai datasource và nhận thông báo kết nối thành công.')

page('Đọc dashboard trạng thái CPU và RAM')
ui(22,'Grafana_CPU_RAM','Dashboard hiển thị sáu target UP cùng CPU và RAM của container',3.8)
table(['Panel hoặc số liệu','Cách hiểu'],[
('Targets đang UP bằng 6','Sáu endpoint metrics được scrape thành công'),
('MySQL kết nối bằng 1','Exporter kết nối được database'),
('Nginx hoạt động bằng 1','Exporter đọc được stub_status'),
('Kết nối MySQL','Tổng kết nối hiện mở gồm cả client giám sát'),
('CPU theo container','Phần trăm công suất một CPU logic theo dịch vụ'),
('RAM theo container','Working set hiển thị theo đơn vị byte MiB GiB')],[2.0,4.5])
p('Đường màu phân biệt app, db, grafana và các dịch vụ khác. Biểu đồ cho thấy hoạt động tăng trong lúc thực hiện truy vấn và chụp minh chứng. Một mẫu tăng chưa đủ kết luận lỗi; cần xem mức kéo dài, giới hạn RAM và log cùng thời điểm.')

page('Giải thích truy vấn CPU RAM và đơn vị đo')
h('CPU trung bình theo dịch vụ')
code('sum by (container_label_com_docker_compose_service) (\n  rate(container_cpu_usage_seconds_total{\n    container_label_com_docker_compose_project="recipe",\n    image!=""\n  }[2m])\n) * 100')
p('Metric CPU là bộ đếm thời gian CPU tích lũy. rate chuyển mức tăng trong cửa sổ hai phút thành giây CPU mỗi giây. Bộ lọc project chỉ chọn Recipe; sum by gom theo tên dịch vụ Compose. Nhân 100 biểu diễn phần trăm của một CPU logic.')
p('Ví dụ dùng 60 giây CPU trong 120 giây tương đương khoảng 50 phần trăm. Trên máy nhiều CPU, giá trị có thể vượt 100 phần trăm. rate xử lý trường hợp counter đặt lại khi tiến trình khởi động lại.')
h('RAM working set theo dịch vụ')
code('sum by (container_label_com_docker_compose_service) (\n  container_memory_working_set_bytes{\n    container_label_com_docker_compose_project="recipe",\n    image!=""\n  }\n)')
p('RAM là giá trị hiện tại nên đọc trực tiếp, không tính rate. Metric trả byte; chia 1024 rồi 1024 để đổi sang MiB. Ví dụ 104857600 byte tương đương 100 MiB. Working set có thể gồm một phần cache, không được đồng nhất với RSS.')
h('Khoảng hiển thị và cửa sổ tính toán')
p('Chọn Last 1 hour xác định chiều dài biểu đồ. Cửa sổ [2m] nằm trong công thức và áp dụng tại từng điểm trên biểu đồ. Vì vậy biểu đồ một giờ vẫn có thể sử dụng tốc độ trung bình hai phút cho mỗi điểm.')

page('Theo dõi Nginx MySQL và lưu lượng mạng')
ui(23,'Grafana_web_database_network','Bốn panel request kết nối MySQL và dữ liệu mạng',3.65)
code('rate(nginx_http_requests_total[2m])\nnginx_connections_active\nmysql_global_status_threads_connected\nrate(mysql_global_status_questions[2m])')
p('Request mỗi giây dùng bộ đếm Nginx; số kết nối Nginx và MySQL là giá trị hiện tại. Questions mỗi giây là tốc độ lệnh từ client gửi tới MySQL, có thể bao gồm phpMyAdmin và exporter. Nếu bộ đếm request tăng 240 trong 120 giây, tốc độ xấp xỉ 2 request mỗi giây.')
code('sum by (container_label_com_docker_compose_service) (\n  rate(container_network_receive_bytes_total{\n    container_label_com_docker_compose_project="recipe",\n    interface!="lo"\n  }[2m])\n)')
p('Truy vấn mạng trả byte nhận mỗi giây, bỏ giao diện loopback. Grafana định dạng thành kB/s hoặc đơn vị phù hợp. Các số đo hạ tầng được dùng để khoanh vùng, không thay thế kiểm thử nghiệp vụ.')

page('Luật cảnh báo Prometheus')
ui(18,'Prometheus_alerts','Ba luật cảnh báo đã nạp và đang ở trạng thái INACTIVE',3.3)
table(['Tên cảnh báo','Biểu thức','Điều kiện thời gian'],[
('RecipeTargetDown','up == 0','Duy trì một phút'),
('RecipeMySQLDown','mysql_up == 0','Duy trì một phút'),
('RecipeNginxDown','nginx_up == 0','Duy trì một phút')],[2.1,2.2,2.2])
code('alert: RecipeMySQLDown\nexpr: mysql_up == 0\nfor: 1m\nlabels:\n  severity: critical')
p('Prometheus đánh giá rule mỗi 15 giây. Khi điều kiện xuất hiện, rule có thể ở trạng thái pending; nếu kéo dài đủ for: 1m mới thành firing. Khi điều kiện hết, cảnh báo trở về trạng thái không kích hoạt.')
p('Ảnh chụp INACTIVE cho thấy tại thời điểm đó các điều kiện lỗi không xảy ra. Cấu hình hiện chưa nối Alertmanager hoặc contact point gửi email; vì vậy không tuyên bố đã có thông báo tự động ra dịch vụ ngoài.')

page('10 Cấu hình Loki Promtail và nguồn log',True)
p('Nginx, PHP và MySQL ghi log vào named volume riêng. Promtail chỉ đọc các volume này, gắn nhãn và gửi log tới Loki. Volume positions giữ vị trí đã đọc. Loki lưu dữ liệu TSDB và chunks trên filesystem trong volume, retention cấu hình 168 giờ.')
code('positions:\n  filename: /positions/positions.yaml\nclients:\n  - url: http://loki:3100/loki/api/v1/push\nscrape_configs:\n  - job_name: nginx-access\n    static_configs:\n      - targets: [localhost]\n        labels:\n          project: recipe\n          job: nginx\n          service: nginx\n          __path__: /logs/nginx/access.log')
table(['Nguồn','Đường dẫn Promtail','Nhãn job'],[
('Nginx access','/logs/nginx/access.log','nginx'),
('Nginx error','/logs/nginx/error.log','nginx_error'),
('Ứng dụng PHP','/logs/app/*.log','app'),
('MySQL error','/logs/mysql/*.log','mysql')],[1.3,3.2,2.0])
p('Prometheus dùng pull để lấy metrics. Promtail dùng push để gửi log. Grafana truy vấn Prometheus bằng PromQL và Loki bằng LogQL. Promtail không được cấp Docker socket; phạm vi đọc log giới hạn ở các volume đã chỉ định.')
p('Promtail được giữ để đáp ứng yêu cầu môn học. Tài liệu Grafana ghi nhận Promtail hết hỗ trợ ngày 02/03/2026; khi phát triển hệ thống vận hành lâu dài cần chuyển sang Alloy hoặc client được hỗ trợ.')

page('Truy vấn LogQL xem access log Nginx')
steps(['Mở Grafana Explore và chọn datasource Loki.','Chuyển sang Code, chọn Last 1 hour và nhập truy vấn dưới đây.','Chạy truy vấn; đối chiếu nhãn nguồn, timestamp và nội dung dòng log.'])
code('{job="nginx"}')
ui(26,'LogQL_Nginx_access','Explore trả các dòng access log và nhãn nguồn Nginx',3.75)
p('Phần trong dấu ngoặc nhọn chọn luồng log có nhãn job=nginx. Mỗi dòng có thời điểm, phương thức HTTP, URI, mã trạng thái, số byte, thời gian xử lý và upstream. Nhãn project=recipe giúp phân biệt dự án khi cùng Loki thu nhiều hệ thống.')
p('Loki tổ chức log theo timestamp và labels. Phân tích JSON trong câu truy vấn diễn ra khi đọc; cấu hình Promtail hiện chỉ gắn nhãn nguồn, không tạo riêng một nhãn chỉ mục cho từng URI hoặc địa chỉ client.')

page('Truy vấn LogQL lọc lỗi HTTP')
code('{job="nginx"} | json | status >= 400')
ui(27,'LogQL_loi_HTTP','Các request lỗi được lọc bằng trường status sau khi parse JSON',3.8)
steps(['Mở một URL món không tồn tại như ?recipe=999999 để tạo HTTP 404.','Truy cập /.env để kiểm tra Nginx từ chối bằng HTTP 403.','Quay lại Explore, chạy truy vấn trên và kiểm tra các dòng trong thời gian vừa thử.'])
p('json tách nội dung log thành trường có tên. status >= 400 chỉ giữ các request có mã từ 400 trở lên. Mã 403 là từ chối truy cập, 404 là không tìm thấy, 500 là lỗi phía máy chủ và 502 thường liên quan trao đổi giữa proxy với upstream.')
p('Ảnh có thể chứa cả 404 của favicon hoặc đường dẫn khác; không suy ra mọi dòng 404 đều do cùng một thao tác thử. Muốn khoanh vùng, đối chiếu timestamp, uri và mã status.')

page('Truy vấn LogQL theo dõi thao tác quản trị')
code('{job="app"} |= "login_success"')
ui(28,'LogQL_dang_nhap','Các sự kiện đăng nhập thành công có timestamp và user_id',3.65)
p('Toán tử |= giữ dòng chứa đoạn văn bản chỉ định. login_success do ứng dụng ghi sau khi xác minh mật khẩu và tạo phiên đăng nhập. Log có tiền tố PHP trước nội dung JSON, nên truy vấn chuỗi là cách trực tiếp để lọc sự kiện trong cấu hình hiện tại.')
table(['Chuỗi lọc','Sự kiện ứng dụng'],[
('login_success','Đăng nhập thành công'),
('login_failed','Đăng nhập thất bại'),
('recipe_saved','Thêm hoặc sửa công thức'),
('recipe_archived','Lưu trữ công thức'),
('recipe_restored','Khôi phục công thức')],[2.0,4.5])
p('Dashboard log dùng {job="app"} |= "event" để xem nhiều loại sự kiện cùng lúc. Các truy vấn theo tên sự kiện dùng khi cần kiểm tra một thao tác cụ thể.')

page('Truy vấn log MySQL và thống kê từ log')
code('{job="mysql"}')
ui(29,'LogQL_MySQL','MySQL error log hiển thị thông báo khởi động và sẵn sàng kết nối',3.45)
p('Nguồn này là error log MySQL, không phải toàn bộ câu SQL đã chạy. Một database khỏe có thể không ghi thêm error log trong nhiều giờ. Vì vậy ảnh chọn Last 7 days để thấy thông báo khởi động đã lưu; test cũng dùng cửa sổ bảy ngày cho nguồn này thay vì bắt buộc phải có log mới trong một giờ.')
h('Đếm request theo mã HTTP trong năm phút')
code('sum by (status) (\n  count_over_time(\n    {job="nginx"} | json | __error__="" [5m]\n  )\n)')
p('count_over_time đếm dòng trong cửa sổ năm phút; sum by status gom theo mã HTTP. Điều kiện __error__="" bỏ dòng lỗi phân tích. Đây là truy vấn sinh số liệu từ log, có thể đưa vào panel biểu đồ. Không có dòng khớp trong khoảng chọn có thể cho kết quả rỗng thay vì số 0.')

page('11 Hardening được kiểm chứng tại runtime',True)
cli('CLI_04_Hardening','Các trường UID read only capabilities và privileged từ Docker inspect',3.55)
cli('CLI_11_UID','Lệnh id xác nhận UID của app Nginx và MySQL',1.65)
p('App và phpMyAdmin dùng UID 33, Nginx 101, MySQL 999, Grafana 472, Prometheus và exporter 65534, Loki và Promtail 10001. Tất cả service bỏ ALL capabilities, bật no-new-privileges và không privileged.')
p('Root filesystem chỉ đọc, trong khi named volume và tmpfs cung cấp các vùng ghi cần thiết. Kiểm tra runtime quan trọng vì file YAML chỉ cho biết cấu hình dự kiến; Docker inspect cho biết các thuộc tính đã thực sự áp dụng cho container.')
p('cAdvisor là ngoại lệ chạy root để đọc cgroups và metadata runtime. Mount host chỉ đọc và không publish cổng làm giảm phạm vi tiếp cận, nhưng socket Unix gắn :ro vẫn cho phép gọi API runtime; collector phải được tin cậy.')

page('Phân quyền database và bảo vệ bí mật')
cli('CLI_08_DDL_bi_tu_choi','MySQL trả mã 1142 khi recipe_user thử tạo bảng',2.25)
table(['Biện pháp','Cách triển khai và kiểm tra'],[
('Quyền tài khoản ứng dụng','Chỉ SELECT INSERT UPDATE DELETE trên recipe_db; SHOW GRANTS và thử CREATE TABLE bị từ chối'),
('Tài khoản giám sát riêng','recipe_exporter có quyền đọc phục vụ metrics và tối đa ba kết nối'),
('Bí mật cục bộ','Mật khẩu ngẫu nhiên, .env mode 0600, secrets mode 0700, git check-ignore xác nhận loại khỏi Git'),
('Tách mạng','db_net và monitoring_net internal; MySQL không publish cổng'),
('Giới hạn tài nguyên','Memory limit từng service; Docker stdout log rotation 10 MB nhân ba file'),
('Cấu hình ổn định','Image bên ngoài ghim digest; Grafana không tự cài và cập nhật plugin khi chạy')],[1.5,5.0],10.1)
p('phpMyAdmin chặn đăng nhập root và sử dụng tài khoản ứng dụng để quản lý các bản ghi. Script setup dùng quyền quản trị để nạp schema trong giai đoạn triển khai; website không giữ quyền tạo bảng.')
p('Ảnh cổng và quyền file ở phần mạng ghi nhận các UI đều bind 127.0.0.1. Báo cáo không hiển thị nội dung .env hoặc hash mật khẩu của bảng users.')

page('12 Kiểm thử chức năng ứng dụng và database',True)
p('Chạy python3 scripts/verify.py sau khi toàn bộ dịch vụ đã khởi động. Script thực hiện request thật, tạo bản ghi tạm, kiểm tra kết quả và dọn bản ghi trong finally. Kết quả được lưu vào minh-chung/verification.json và verification.txt.')
cli('CLI_14_Test_ung_dung','Nhóm kiểm thử HTTP tìm kiếm CRUD xác thực và database',4.8)
p('Chu trình CRUD kiểm tra cả giao diện công khai lẫn thay đổi trạng thái trong database. Tình huống bảo mật thử CSRF, mật khẩu sai, SQL injection và HTML escape. Kiểm tra phpMyAdmin đăng nhập bằng tài khoản ứng dụng và đọc được năm bảng chính.')
p('Kết quả hiện tại là '+str(V['passed'])+'/'+str(V['total'])+' đạt tại '+CHECK+'. Đây là phạm vi kiểm thử của bản lab, không thay thế kiểm toán bảo mật hoặc thử tải cho một hệ thống vận hành lớn.')

page('Kiểm thử giám sát log và hardening')
cli('CLI_15_Test_monitoring','Kết quả kiểm thử target metrics datasource dashboard và bốn LogQL',3.7)
cli('CLI_16_Test_runtime','Kết quả kiểm thử cấu hình runtime và cổng mạng',3.7)
p('Bộ kiểm thử xác nhận sáu target UP, metrics CPU/RAM của mười một container, hai datasource healthy, dashboard metrics có mười panel và bốn LogQL trả dữ liệu. Các kiểm tra runtime xác nhận UID theo vai trò, read_only, cap_drop ALL, no-new-privileges, loopback và db_net internal.')

page('Kiểm tra cài mới và sao lưu database')
cli('CLI_17_Cai_moi','Tám kiểm tra cài mới độc lập từ GitHub ở commit đã ghi nhận',3.6)
cli('CLI_18_Backup','Bản sao lưu mới có quyền 0600 và đọc được định dạng gzip',2.1)
p('Kiểm tra cài mới thực hiện ngày 23/10/2026 tại commit '+R['commit'][:7]+', project recipe-repro, volume và cổng khác hệ thống chính. Kết quả có đúng năm bảng, sáu công thức tiếng Việt, admin đăng nhập được, hai datasource hoạt động và đủ dịch vụ. Các container thử được dọn sau kiểm tra.')
p('Bản sao lưu mới được tạo trong lần bổ sung báo cáo bằng scripts/backup.py. File gzip nằm trong runtime/backups, được Git bỏ qua. Kiểm tra đọc gzip xác nhận file có dữ liệu; chưa thực hiện phục hồi ghi đè database đang chạy trong đợt cập nhật báo cáo này.')

page('13 Vận hành và các vấn đề đã xử lý',True)
code('docker compose up -d\ndocker compose ps -a\ndocker compose logs --tail=50 app nginx\npython3 scripts/backup.py\ndocker compose stop')
table(['Vấn đề gặp khi triển khai','Nguyên nhân và cách xử lý đã áp dụng'],[
('Container chưa áp dụng cổng mới','Tạo lại đúng service để Docker áp dụng PortBindings từ Compose'),
('Ứng dụng cũ theo WordPress','Chuyển sang PHP theo khung mẫu; lưu bản sao và giữ bảng volume cũ'),
('Ghi file trên root filesystem chỉ đọc','Chuẩn bị thư mục và quyền trong image; dùng volume hoặc tmpfs cho dữ liệu ghi'),
('Nginx giữ địa chỉ app cũ','Dùng Docker resolver và upstream resolve để cập nhật DNS'),
('Grafana cập nhật plugin khi runtime chỉ đọc','Tắt tự preinstall và update, giữ datasource được provision'),
('Grafana thiếu RAM khi mở dashboard','Tăng giới hạn từ 384 MiB lên 768 MiB, đặt GOMEMLIMIT và giảm plugin không dùng'),
('MySQL error log không có dòng trong một giờ','Database không phát sinh lỗi mới; chọn cửa sổ bảy ngày và sửa false failure của test')],[2.2,4.3],10.1)
p('Những mục trên mô tả vấn đề đã gặp trong quá trình thực hiện và cấu hình hiện đã áp dụng. Ảnh trạng thái, headers, dashboard và các trường runtime ở các phần trước chứng minh kết quả sau xử lý; không dựng lại ảnh lỗi cũ để thay bằng chứng.')
p('Khi gặp No data, kiểm tra datasource, khoảng thời gian, nhãn và thời điểm phát sinh sự kiện trước khi kết luận dịch vụ hỏng. Với biểu đồ rate, cần có đủ mẫu scrape trong cửa sổ tính.')

page('Kịch bản demo giới hạn và tài liệu tham khảo')
steps(['Mở GitHub, giới thiệu source và ba giai đoạn commit kỹ thuật.','Chạy docker compose ps, mở website, tìm nguyên liệu và lọc danh mục.','Xem chi tiết món rồi đăng nhập quản trị, minh họa thêm sửa lưu trữ khôi phục bằng dữ liệu riêng.','Mở phpMyAdmin đối chiếu sáu công thức bằng SELECT JOIN.','Mở Prometheus Targets và Grafana, giải thích sáu target và mười một container.','Chạy các LogQL Nginx, lỗi HTTP, login_success và MySQL; giải thích time range.','Chỉ ra hardening từ Docker inspect, quyền DB và cổng loopback.','Trình bày kết quả verify, sao lưu và cách dừng mà vẫn giữ volume.'])
p('Giới hạn của bản lab: HTTP localhost với security headers; chưa có HTTPS công khai, thông báo cảnh báo ngoài hệ thống, lịch xoay log file nguồn, high availability hay đánh giá tải lớn. Dữ liệu mẫu gồm sáu công thức. Promtail cần được thay bằng client được hỗ trợ nếu phát triển vận hành lâu dài.')
h('Tài liệu tham khảo chính thức và source')
refs=[
('Mã nguồn Recipe',REPO),
('Khung thư mục Blog CMS','https://github.com/Khanhslee/DTC245220024_De7-Blog-CMS'),
('Prometheus tổng quan','https://prometheus.io/docs/introduction/overview/'),
('Hàm PromQL','https://prometheus.io/docs/prometheus/latest/querying/functions/'),
('Grafana datasource Prometheus','https://grafana.com/docs/grafana/latest/datasources/prometheus/'),
('Loki LogQL','https://grafana.com/docs/loki/latest/query/log_queries/'),
('Promtail và vòng đời hỗ trợ','https://grafana.com/docs/loki/latest/send-data/promtail/'),
('Nginx stub status','https://nginx.org/en/docs/http/ngx_http_stub_status_module.html')]
for title,url in refs:
    x=doc.add_paragraph(title+'\n'+url);x.paragraph_format.line_spacing=1.0;x.paragraph_format.space_after=Pt(4)
    for r in x.runs:r.font.size=Pt(9)

page('Phụ lục danh mục hình và nguồn minh chứng',True)
p('Danh mục giúp đối chiếu số hình trong báo cáo với file ảnh trong repository. UI là ảnh chụp giao diện thật; CLI là hình trích output lệnh có TXT và JSON nguồn. Ảnh sơ đồ thể hiện cấu trúc thiết kế từ cấu hình.')
table(['Hình','Nội dung minh chứng','Trang'],[(f['number'],f['caption'],f['planned_page']) for f in figures],[.45,5.35,.7],9.5)
p('Các file nguồn của hình CLI nằm trong minh-chung/chi-tiet, cùng tên với PNG và phần mở rộng TXT. command-evidence.json giữ command, timestamp, exit code và output. verification.json giữ từng kết quả kiểm thử; reproducibility.json giữ mốc cài mới độc lập. Mật khẩu và bản dump database không nằm trong phần phụ lục công khai.')

# Mục lục tĩnh được tính từ các ngắt trang chủ động. PDF QA đối chiếu số trang thực.
for title,num in chapter_rows:
    cells=toc.add_row().cells
    cells[0].text=title;cells[1].text=str(num)
    for c in cells:
        for q in c.paragraphs:
            q.paragraph_format.line_spacing=1.05;q.paragraph_format.space_after=Pt(3)
            for rr in q.runs:rr.font.size=Pt(10.5)
    cells[1].paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
    for cell in cells:
        borders=OxmlElement('w:tcBorders')
        for edge in ['top','left','bottom','right']:
            q=OxmlElement('w:'+edge);q.set(qn('w:val'),'single');q.set(qn('w:sz'),'4');q.set(qn('w:color'),'D9D9D9');borders.append(q)
        cell._tc.get_or_add_tcPr().append(borders)
for tree in [doc.styles.element,doc._element]:
    for border in list(tree.iter(qn('w:pBdr'))):border.getparent().remove(border)
path=OUT/'Bao_cao_De29_Recipe.docx';doc.save(path)
(OUT/'qa').mkdir(exist_ok=True)
(OUT/'qa/report-manifest.json').write_text(json.dumps({'planned_pages':pages,'figures':figures,'chapter_rows':chapter_rows,'verification_time':V['time'],'capture_time':C['time']},ensure_ascii=False,indent=2))
print(json.dumps({'path':str(path),'planned_content_pages':len(pages),'figures':len(figures)},ensure_ascii=False))
